import docx,json,re,glob
files={'lab':'LABORATORIO CLINICO','rad':'RADIOLOGIA','ter':'TERAPIA FÍSCIA'}
out={}
for k,n in files.items():
    d=docx.Document(f'../demanda/Demanda Social - {n}.docx')
    out[k]=[[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables]
    print(k,len(d.tables),[ (len(t.rows),len(t.columns)) for t in d.tables])
json.dump(out,open('tables.json','w'),ensure_ascii=False)
